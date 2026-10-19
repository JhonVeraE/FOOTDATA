"""
Genera un Excel por cada partido de jornada_predictions.json, con el
nuevo nivel de detalle (multi-ventana, tendencias, cruces ataque-defensa,
H2H, confianza) -- reemplaza al build_excel_jornada1.py anterior, que
solo mostraba lineas de mercado.

Uso:
    python3 build_excel_avanzado.py                  # un Excel por partido
    python3 build_excel_avanzado.py --liga "La Liga" --jornada 1
"""
import argparse
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from paths import DATA_DIR
from paths import ARTIFACT_DIR
OUT_DIR = ARTIFACT_DIR / "excel_partidos"

FONT_NAME = "Calibri"
NAVY = "0D1428"
GOLD = "F5B841"
BLUE = "3B82F6"
RED = "EF4444"
GRAY = "E5E7EB"
GREEN = "16A34A"


def _title(cell, size=18):
    cell.font = Font(name=FONT_NAME, bold=True, size=size, color=NAVY)
    cell.alignment = Alignment(horizontal="center")


def _header(cell, size=12):
    cell.font = Font(name=FONT_NAME, bold=True, size=size, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _sub_header(cell, size=10):
    cell.font = Font(name=FONT_NAME, bold=True, size=size, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color="30415E")
    cell.alignment = Alignment(horizontal="center")


def build_match_excel(m: dict, out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [24, 16, 16, 16, 16, 16, 16]):
        ws.column_dimensions[col].width = w

    home, away = m["local"], m["visitante"]

    r = 1
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = f"{home}  vs  {away}"
    _title(ws[f"A{r}"])
    r += 1
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = f"{m['fecha']} {m.get('hora','')} UTC · {m['competition']} · Jornada {m.get('jornada','?')}"
    ws[f"A{r}"].font = Font(name=FONT_NAME, size=11, italic=True, color="6B7280")
    ws[f"A{r}"].alignment = Alignment(horizontal="center")

    # ---- 1X2 + confianza ----
    r += 2
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "PROBABILIDAD DE RESULTADO (1X2)"
    _header(ws[f"A{r}"])
    r += 1
    labels = [f"Gana {home}", "Empate", f"Gana {away}"]
    values = [m["p_home"], m["p_draw"], m["p_away"]]
    colors = [BLUE, "9CA3AF", RED]
    for i, lab in enumerate(labels):
        c = ws.cell(row=r, column=1 + i * 2)
        ws.merge_cells(start_row=r, start_column=1 + i * 2, end_row=r, end_column=2 + i * 2)
        c.value = lab; c.font = Font(name=FONT_NAME, size=11, bold=True); c.alignment = Alignment(horizontal="center")
    r += 1
    for i, (val, col) in enumerate(zip(values, colors)):
        c = ws.cell(row=r, column=1 + i * 2)
        ws.merge_cells(start_row=r, start_column=1 + i * 2, end_row=r, end_column=2 + i * 2)
        c.value = f"{val:.1f}%"; c.font = Font(name=FONT_NAME, size=22, bold=True, color=col)
        c.alignment = Alignment(horizontal="center")
    r += 2
    conf = m["confianza"]
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = f"Calidad de muestra: {conf['nivel']}  ·  Goles esperados: {home} {m['lam_home']:.2f} — {away} {m['lam_away']:.2f}"
    ws[f"A{r}"].font = Font(name=FONT_NAME, size=11, italic=True)
    ws[f"A{r}"].alignment = Alignment(horizontal="center")

    # ---- Marcadores ----
    r += 2
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "MARCADORES MÁS PROBABLES (distribución de goles)"
    _header(ws[f"A{r}"])
    r += 1
    ws.cell(row=r, column=1, value="Marcador").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=r, column=2, value="Probabilidad").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=r, column=4, value="Ambos anotan (BTTS)").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=r, column=5, value=f"{m['btts']}%").font = Font(name=FONT_NAME, bold=True, color=GOLD)
    for score, prob in m["marcadores_top5"]:
        r += 1
        ws.cell(row=r, column=1, value=score).font = Font(name=FONT_NAME, size=12)
        ws.cell(row=r, column=2, value=f"{prob}%").font = Font(name=FONT_NAME, size=12)

    # ---- Forma reciente multi-ventana (seccion 1) ----
    r += 2
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "FORMA RECIENTE — TEMPORADA vs ÚLTIMOS 10 vs ÚLTIMOS 5 vs ÚLTIMOS 3"
    _header(ws[f"A{r}"])
    r += 1

    def tabla_forma(equipo_nombre, forma, ctx_label):
        nonlocal r
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = f"{equipo_nombre} ({ctx_label})"
        _sub_header(ws[f"A{r}"])
        r += 1
        cols = [("Ventana", None), ("Partidos", "partidos"), ("Pts/PJ", "puntos_por_partido"),
                ("Goles favor", "goals_for"), ("Goles contra", "goals_against"),
                ("Corners favor", "corners_for"), ("Tarjetas favor", "cards_for")]
        for j, (h, _) in enumerate(cols):
            c = ws.cell(row=r, column=1 + j, value=h)
            c.font = Font(name=FONT_NAME, bold=True, size=9)
            c.fill = PatternFill("solid", start_color=GRAY)
        r += 1
        for ventana_key, ventana_label in [("temporada", "Temporada"), ("ultimos_10", "Últimos 10"),
                                             ("ultimos_5", "Últimos 5"), ("ultimos_3", "Últimos 3")]:
            if ventana_key not in forma:
                continue
            v = forma[ventana_key]
            vals = [ventana_label, v["partidos"], v["puntos_por_partido"], v["goals_for"],
                    v["goals_against"], v["corners_for"], v["cards_for"]]
            for j, val in enumerate(vals):
                ws.cell(row=r, column=1 + j, value=val).font = Font(name=FONT_NAME, size=10)
            r += 1
        r += 1

    tabla_forma(home, m["forma_home"], "local")
    tabla_forma(away, m["forma_away"], "visitante")

    # ---- Tendencias detectadas ----
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "TENDENCIAS DETECTADAS"
    _header(ws[f"A{r}"])
    r += 1
    for equipo_nombre, tendencias in [(home, m["tendencias_home"]), (away, m["tendencias_away"])]:
        cambios = [(k, v) for k, v in tendencias.items() if v["estado"] in ("subiendo", "bajando")]
        if not cambios:
            ws.cell(row=r, column=1, value=f"{equipo_nombre}: sin tendencia clara detectada (estable)").font = Font(name=FONT_NAME, size=10, italic=True)
            r += 1
            continue
        for stat, td in cambios:
            color = BLUE  # Direction alone is not a sporting improvement/deterioration.
            texto = f"{equipo_nombre}: {td['estado'].upper()} en {stat.replace('_',' ')} ({td['cambio_pct_ultimo']:+.0f}%)"
            ws.cell(row=r, column=1, value=texto).font = Font(name=FONT_NAME, size=10, color=color)
            r += 1
    r += 1
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "Nota: los cambios de 'últimos 3' pueden reflejar ruido estadístico (muestra chica), no solo una tendencia real."
    ws[f"A{r}"].font = Font(name=FONT_NAME, size=9, italic=True, color="9CA3AF")
    r += 2

    # ---- Cruces ataque-defensa (seccion 6) ----
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "CRUCES ATAQUE vs DEFENSA (con forma reciente ponderada)"
    _header(ws[f"A{r}"])
    r += 1
    headers = ["Estadística", f"Proyección {home}", "",
               f"Proyección {away}", "", "Total esperado", ""]
    for j, h in enumerate(headers):
        c = ws.cell(row=r, column=1 + j, value=h)
        c.font = Font(name=FONT_NAME, bold=True, size=9)
        c.fill = PatternFill("solid", start_color=GRAY)
    r += 1
    cc = m["cruces"]
    nombres = {"goals": "Goles", "shots": "Remates", "sot": "Remates al arco",
               "corners": "Corners", "cards": "Tarjetas", "fouls": "Faltas"}
    for stat, label in nombres.items():
        v = cc[stat]
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT_NAME, size=10, bold=True)
        ws.cell(row=r, column=2, value=v["home"]).font = Font(name=FONT_NAME, size=10, color=BLUE)
        ws.cell(row=r, column=4, value=v["away"]).font = Font(name=FONT_NAME, size=10, color=RED)
        total=round(v["home"]+v["away"],2) if v['home'] is not None and v['away'] is not None else 'sin datos'
        ws.cell(row=r, column=6, value=total).font = Font(name=FONT_NAME, size=10, bold=True)
        r += 1
    r += 1
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = ("Metodología: ataque del local (de local, ponderado por recencia) combinado con defensa del "
                    "visitante (de visita), y viceversa. Goles: ataque/defensa ajustados por rival. H2H descriptivo, sin ajuste de goles.")
    ws[f"A{r}"].font = Font(name=FONT_NAME, size=9, italic=True, color="6B7280")
    ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    r += 2

    # ---- H2H ----
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = f"ENFRENTAMIENTOS DIRECTOS (H2H) — {m.get('h2h_total_partidos', 0)} en el historial cargado"
    _header(ws[f"A{r}"])
    r += 1
    h2h_ida = m.get("h2h_ida", {})
    detalle = h2h_ida.get("detalle", []) if h2h_ida else []
    if detalle:
        headers = ["Fecha", "Local", "Resultado", "Corners", "Tarjetas"]
        for j, h in enumerate(headers):
            c = ws.cell(row=r, column=1 + j, value=h)
            c.font = Font(name=FONT_NAME, bold=True, size=9)
            c.fill = PatternFill("solid", start_color=GRAY)
        r += 1
        for d in detalle:
            ws.cell(row=r, column=1, value=d["fecha"]).font = Font(name=FONT_NAME, size=10)
            ws.cell(row=r, column=2, value=home).font = Font(name=FONT_NAME, size=10)
            ws.cell(row=r, column=3, value=d["marcador"]).font = Font(name=FONT_NAME, size=10, bold=True)
            ws.cell(row=r, column=4, value=d["corners"]).font = Font(name=FONT_NAME, size=10)
            ws.cell(row=r, column=5, value=d["tarjetas"]).font = Font(name=FONT_NAME, size=10)
            r += 1
    else:
        ws.cell(row=r, column=1, value="Sin enfrentamientos directos en el historial cargado.").font = Font(name=FONT_NAME, italic=True, color="9CA3AF")
        r += 1
    r += 1

    # ---- Conclusión ----
    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "CONCLUSIÓN DEL MODELO"
    _header(ws[f"A{r}"])
    r += 1
    ws.merge_cells(f"A{r}:G{r+1}")
    ws[f"A{r}"] = m["conclusion"]
    ws[f"A{r}"].font = Font(name=FONT_NAME, size=11)
    ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    r += 3

    # ---- Datos para contenido ----
    if m.get("datos_curiosos"):
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"] = "DATOS DESTACADOS (para redes/contenido)"
        _header(ws[f"A{r}"])
        r += 1
        for d in m["datos_curiosos"]:
            ws.merge_cells(f"A{r}:G{r}")
            ws[f"A{r}"] = f"• {d}"
            ws[f"A{r}"].font = Font(name=FONT_NAME, size=10, color=GOLD)
            ws[f"A{r}"].alignment = Alignment(wrap_text=True)
            r += 1
        r += 1

    ws.merge_cells(f"A{r}:G{r}")
    ws[f"A{r}"] = "Ataque/defensa regularizados con corte temporal. Córners/tarjetas: estimaciones sin calibrar. No garantiza resultados."
    ws[f"A{r}"].font = Font(name=FONT_NAME, size=9, italic=True, color="9CA3AF")
    ws[f"A{r}"].alignment = Alignment(horizontal="center", wrap_text=True)

    wb.save(out_path)


def main(liga_filtro=None, jornada_filtro=None):
    json_path = os.path.join(DATA_DIR, "jornada_predictions.json")
    if not os.path.exists(json_path):
        raise SystemExit("Falta jornada_predictions.json. Corre antes: python3 generate_jornada_predictions.py")
    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)

    os.makedirs(OUT_DIR, exist_ok=True)
    n = 0
    for key, m in data.items():
        if liga_filtro and m["competition"] != liga_filtro:
            continue
        if jornada_filtro and m.get("jornada") != jornada_filtro:
            continue
        fname = f"Prediccion_{m['local']}_vs_{m['visitante']}.xlsx".replace(" ", "_").replace("/", "-")
        out_path = os.path.join(OUT_DIR, fname)
        build_match_excel(m, out_path)
        print(f"  OK -> {fname}")
        n += 1
    print(f"\nListo. {n} archivos generados en {OUT_DIR}/")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--liga", type=str, default=None)
    ap.add_argument("--jornada", type=int, default=None)
    args = ap.parse_args()
    main(liga_filtro=args.liga, jornada_filtro=args.jornada)
